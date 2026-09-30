"""Фильтры, снимки ответов и выгрузка в Excel."""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from maxtest.core.enums import FinishReason, QuestionType
from maxtest.core.grader import grade_attempt
from maxtest.core.models import Option, Question, Scale, Section, Test
from maxtest.core.session import build_plan
from maxtest.core.storage import Storage
from maxtest.report.excel import export_attempts, stats_of


@pytest.fixture()
def storage(tmp_path, monkeypatch):
    monkeypatch.setenv("MAXTEST_DATA_DIR", str(tmp_path))
    with Storage(tmp_path / "results.db") as st:
        yield st


def make_test(title: str = "Аттестация 2026") -> Test:
    test = Test(
        title=title,
        sections=[Section(id="s1", title="Охрана труда"), Section(id="s2", title="Оборудование")],
    )
    test.settings.shuffle_questions = False
    test.grading.scale = Scale.default_pass_fail(70)
    test.questions = [
        Question(
            id=f"q{i}",
            type=QuestionType.SINGLE,
            text=f"Вопрос {i} — что делать?",
            section_id="s1" if i < 2 else "s2",
            explanation=f"Пояснение к вопросу {i}",
            options=[
                Option(id=f"a{i}", text=f"Правильно {i}", correct=True),
                Option(id=f"b{i}", text=f"Неправильно {i}"),
            ],
        )
        for i in range(4)
    ]
    return test


def save_attempt(
    storage: Storage,
    name: str,
    right: int = 4,
    test: Test | None = None,
    started_at: str | None = None,
) -> int:
    test = test or make_test()
    plan = build_plan(test, seed=1)
    answers = {
        f"q{i}": {"option_id": f"a{i}" if i < right else f"b{i}"} for i in range(4)
    }
    result = grade_attempt(test, plan, answers)
    employee_id = storage.get_or_create_employee(name)
    return storage.save_attempt(
        test, plan, answers, result,
        employee_name=name, employee_id=employee_id, started_at=started_at,
    )


# ------------------------------------------------------------------- фильтры


def test_filter_by_employee(storage):
    save_attempt(storage, "Иванов И.И.")
    save_attempt(storage, "Петров П.П.")
    ivanov = storage.find_employee("Иванов И.И.")

    rows = storage.list_attempts(employee_id=ivanov.id)
    assert len(rows) == 1 and rows[0]["employee_name"] == "Иванов И.И."


def test_filter_by_test(storage):
    save_attempt(storage, "Иванов И.И.")
    other = make_test("Пожарная безопасность")
    save_attempt(storage, "Иванов И.И.", test=other)

    tests = storage.list_tests()
    assert len(tests) == 2
    target = next(t for t in tests if t["test_title"] == "Пожарная безопасность")
    rows = storage.list_attempts(test_id=target["test_id"])
    assert len(rows) == 1


def test_filter_by_date_range_is_inclusive(storage):
    save_attempt(storage, "Иванов И.И.", started_at="2026-03-10T09:00:00")
    save_attempt(storage, "Петров П.П.", started_at="2026-03-15T18:30:00")

    rows = storage.list_attempts(date_from="2026-03-10", date_to="2026-03-15")
    assert len(rows) == 2  # обе границы входят, время в конце дня не отсекается

    rows = storage.list_attempts(date_from="2026-03-11", date_to="2026-03-14")
    assert rows == []


def test_filter_only_failed(storage):
    save_attempt(storage, "Иванов И.И.", right=4)
    save_attempt(storage, "Петров П.П.", right=1)

    rows = storage.list_attempts(only_failed=True)
    assert len(rows) == 1 and rows[0]["employee_name"] == "Петров П.П."


def test_delete_attempt_removes_answers(storage):
    attempt_id = save_attempt(storage, "Иванов И.И.")
    storage.delete_attempt(attempt_id)
    assert storage.list_attempts() == []
    assert storage.get_answers(attempt_id) == []


# -------------------------------------------------------------------- снимки


def test_answer_snapshots_are_self_sufficient(storage):
    """Разбор не должен требовать файла .qtest."""
    attempt_id = save_attempt(storage, "Иванов И.И.", right=1)
    answers = {a["question_id"]: a for a in storage.get_answers(attempt_id)}

    assert answers["q0"]["given_answer"] == "Правильно 0"
    assert answers["q3"]["given_answer"] == "Неправильно 3"
    assert answers["q3"]["correct_answer"] == "Правильно 3"
    assert answers["q3"]["explanation"] == "Пояснение к вопросу 3"
    assert answers["q0"]["section_title"] == "Охрана труда"


def test_section_totals(storage):
    attempt_id = save_attempt(storage, "Иванов И.И.", right=2)
    totals = {r["section"]: r for r in storage.section_totals(attempt_id)}

    assert totals["Охрана труда"]["score"] == pytest.approx(2.0)
    assert totals["Оборудование"]["score"] == pytest.approx(0.0)
    assert totals["Оборудование"]["questions"] == 2


def test_hardest_questions_sorted_by_error_rate(storage):
    save_attempt(storage, "Иванов И.И.", right=2)
    save_attempt(storage, "Петров П.П.", right=2)
    ids = [row["id"] for row in storage.list_attempts()]

    hardest = storage.hardest_questions(ids)
    assert hardest[0]["correct"] == 0  # вопросы, где ошиблись все, — первыми
    assert hardest[-1]["correct"] == hardest[-1]["asked"]


# --------------------------------------------------------------------- Excel


def test_export_creates_four_sheets(storage, tmp_path):
    save_attempt(storage, "Иванов И.И.", right=4)
    save_attempt(storage, "Петров П.П.", right=1)

    path = export_attempts(storage, storage.list_attempts(), tmp_path / "отчёт.xlsx")
    book = load_workbook(path)
    assert book.sheetnames == ["Сводная", "По вопросам", "По разделам", "Сложные вопросы"]


def test_export_adds_extension(storage, tmp_path):
    save_attempt(storage, "Иванов И.И.")
    path = export_attempts(storage, storage.list_attempts(), tmp_path / "отчёт")
    assert path.suffix == ".xlsx"


def test_export_keeps_cyrillic(storage, tmp_path):
    save_attempt(storage, "Пётр Фёдорович", right=3)
    path = export_attempts(storage, storage.list_attempts(), tmp_path / "r.xlsx")

    sheet = load_workbook(path)["Сводная"]
    values = [cell.value for row in sheet.iter_rows() for cell in row]
    assert "Пётр Фёдорович" in values
    assert "Аттестация 2026" in values


def test_export_details_contain_answers(storage, tmp_path):
    save_attempt(storage, "Иванов И.И.", right=1)
    path = export_attempts(storage, storage.list_attempts(), tmp_path / "r.xlsx")

    sheet = load_workbook(path)["По вопросам"]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 4
    assert any("Неправильно 3" in str(row) for row in rows)
    assert any("Пояснение" not in str(row) for row in rows)  # пояснения — в разборе


def test_export_sections_percent(storage, tmp_path):
    save_attempt(storage, "Иванов И.И.", right=2)
    path = export_attempts(storage, storage.list_attempts(), tmp_path / "r.xlsx")

    sheet = load_workbook(path)["По разделам"]
    rows = {row[0]: row for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert rows["Охрана труда"][3] == 100
    assert rows["Оборудование"][3] == 0


def test_export_column_widths_are_set(storage, tmp_path):
    """Без ручного расчёта ширины отчёт нечитаем: openpyxl не умеет auto-fit."""
    save_attempt(storage, "Иванов И.И.")
    path = export_attempts(storage, storage.list_attempts(), tmp_path / "r.xlsx")

    sheet = load_workbook(path)["По вопросам"]
    widths = [dim.width for dim in sheet.column_dimensions.values()]
    assert widths and all(w >= 9 for w in widths)


def test_export_empty_selection_still_valid(storage, tmp_path):
    path = export_attempts(storage, [], tmp_path / "пусто.xlsx")
    assert load_workbook(path)["Сводная"]["A1"].value == "Результаты тестирования"


def test_stats_of():
    class Row(dict):
        def __getitem__(self, key):
            return dict.__getitem__(self, key)

    rows = [Row(passed=1, percent=80), Row(passed=0, percent=40), Row(passed=1, percent=90)]
    stats = stats_of(rows)
    assert (stats.attempts, stats.passed, stats.failed) == (3, 2, 1)
    assert stats.average_percent == 70.0


def test_note_column_marks_timeout(storage, tmp_path):
    test = make_test()
    plan = build_plan(test, seed=1)
    answers = {"q0": {"option_id": "a0"}}
    result = grade_attempt(test, plan, answers)
    storage.save_attempt(
        test, plan, answers, result,
        employee_name="Иванов", finish_reason=FinishReason.TIMEOUT,
    )

    path = export_attempts(storage, storage.list_attempts(), tmp_path / "r.xlsx")
    sheet = load_workbook(path)["Сводная"]
    values = [cell.value for row in sheet.iter_rows() for cell in row]
    assert "время вышло" in values
