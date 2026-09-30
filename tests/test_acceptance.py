"""Приёмочный прогон: реальный сценарий заказчика от начала до конца.

Двадцать вопросов, три сотрудника, выборка по разделам, перемешивание,
ручная проверка, выгрузка в Excel. Все ожидаемые числа посчитаны на бумаге и
записаны в тест — если движок начнёт считать иначе, это будет видно сразу.
"""

from __future__ import annotations

import os

import pytest
from openpyxl import load_workbook

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from maxtest.core.bundle import Bundle  # noqa: E402
from maxtest.core.enums import (  # noqa: E402
    MultiMode,
    QuestionType,
    SelectionMode,
    TextMatch,
)
from maxtest.core.grader import grade_attempt  # noqa: E402
from maxtest.core.models import (  # noqa: E402
    MatchPair,
    Option,
    OrderItem,
    Question,
    Scale,
    Section,
    TextAnswer,
)
from maxtest.core.session import build_plan  # noqa: E402
from maxtest.core.storage import Storage  # noqa: E402
from maxtest.core.validator import errors, validate  # noqa: E402
from maxtest.report.excel import export_attempts  # noqa: E402


@pytest.fixture()
def data_dir(tmp_path, monkeypatch):
    monkeypatch.setenv("MAXTEST_DATA_DIR", str(tmp_path))
    return tmp_path


def build_annual_test(tmp_path) -> Bundle:
    """Годовая аттестация: 20 вопросов, 2 раздела, все пять типов."""
    bundle = Bundle.new("Годовая аттестация электромонтёров 2026")
    test = bundle.test
    test.author = "Отдел охраны труда"
    test.sections = [
        Section(id="s_ot", title="Охрана труда", take_count=5),
        Section(id="s_eq", title="Оборудование", take_count=5),
    ]
    test.settings.selection_mode = SelectionMode.BY_SECTION
    test.settings.shuffle_questions = True
    test.settings.shuffle_options = True
    test.settings.time_limit_sec = 1800
    test.grading.multi_mode = MultiMode.PARTIAL
    test.grading.text_manual_review = True
    test.grading.scale = Scale.default_five_point()

    questions: list[Question] = []
    for index in range(8):  # 8 вопросов «один ответ» в разделе охраны труда
        questions.append(
            Question(
                id=f"ot{index}",
                type=QuestionType.SINGLE,
                section_id="s_ot",
                text=f"ОТ-{index}: как поступить в ситуации {index}?",
                explanation="ПОТЭЭ, п. 3.2",
                options=[
                    Option(id=f"ot{index}a", text="Обесточить участок", correct=True),
                    Option(id=f"ot{index}b", text="Продолжить работу"),
                    Option(id=f"ot{index}c", text="Все перечисленное", no_shuffle=True),
                ],
            )
        )
    for index in range(8):  # 8 вопросов в разделе оборудования
        questions.append(
            Question(
                id=f"eq{index}",
                type=QuestionType.MULTI,
                section_id="s_eq",
                weight=2.0,
                text=f"ОБ-{index}: какие средства защиты применяются?",
                options=[
                    Option(id=f"eq{index}a", text="Перчатки", correct=True),
                    Option(id=f"eq{index}b", text="Коврик", correct=True),
                    Option(id=f"eq{index}c", text="Отвёртка"),
                    Option(id=f"eq{index}d", text="Молоток"),
                ],
            )
        )

    questions.append(
        Question(
            id="txt1",
            type=QuestionType.TEXT,
            text="Опишите порядок допуска к работам",
            answer_text=TextAnswer(accepted=["наряд-допуск"], match=TextMatch.CONTAINS),
        )
    )
    questions.append(
        Question(
            id="match1",
            type=QuestionType.MATCHING,
            text="Сопоставьте прибор и величину",
            pairs=[
                MatchPair(id="p1", left="Амперметр", right="Ток"),
                MatchPair(id="p2", left="Вольтметр", right="Напряжение"),
            ],
        )
    )
    questions.append(
        Question(
            id="ord1",
            type=QuestionType.ORDERING,
            text="Расставьте мероприятия по порядку",
            order=[
                OrderItem(id="i1", text="Отключение"),
                OrderItem(id="i2", text="Плакаты"),
                OrderItem(id="i3", text="Проверка напряжения"),
                OrderItem(id="i4", text="Заземление"),
            ],
        )
    )
    questions.append(
        Question(
            id="single_free",
            type=QuestionType.SINGLE,
            text="Итоговый вопрос вне разделов",
            options=[
                Option(id="f_a", text="Верно", correct=True),
                Option(id="f_b", text="Неверно"),
            ],
        )
    )

    test.questions = questions
    bundle.save(tmp_path / "аттестация.qtest", bump_revision=False)
    return bundle


# --------------------------------------------------------------------------


def test_annual_attestation_end_to_end(data_dir, tmp_path):
    bundle = build_annual_test(tmp_path)

    # 1. Тест проходит проверку без ошибок.
    assert errors(validate(bundle.test, known_media=set(bundle.media))) == []
    assert len(bundle.test.questions) == 20

    # 2. Файл переносится: перечитываем с диска, как на рабочем ПК.
    bundle = Bundle.load(tmp_path / "аттестация.qtest")

    with Storage(tmp_path / "results.db") as storage:
        plans = {}
        for name in ("Иванов И.И.", "Петров П.П.", "Сидорова А.А."):
            employee_id = storage.get_or_create_employee(name)
            plan = build_plan(bundle.test, seed=hash(name) % 10_000)
            plans[name] = plan

            # Отбор по разделам: 5 + 5 из разделов плюс 4 вопроса вне разделов.
            assert len(plan.question_ids) == 14

            answers = _answers_for(bundle, plan, name)
            result = grade_attempt(bundle.test, plan, answers)
            storage.save_attempt(
                bundle.test, plan, answers, result,
                employee_name=name, employee_id=employee_id,
            )

        # 3. Три попытки, ни одна не затёрла другую.
        attempts = storage.list_attempts()
        assert len(attempts) == 3
        assert {a["employee_name"] for a in attempts} == {
            "Иванов И.И.", "Петров П.П.", "Сидорова А.А."
        }

        # 4. Максимум одинаков у всех: 5 одиночных + 5 двойных + 4 прочих.
        #    5*1 + 5*2 + (1 текст + 1 соответствие + 1 порядок + 1 одиночный) = 19
        assert {a["max_score"] for a in attempts} == {19.0}

        # 5. Все текстовые ответы ушли на ручную проверку.
        assert storage.count_attempts_to_review() == 3

        # 6а. Иванов ответил верно, движок уже начислил балл: подтверждение
        #     проверяющего снимает флаг, но итог не меняет.
        ivanov = next(a for a in attempts if a["employee_name"] == "Иванов И.И.")
        before = storage.get_attempt(ivanov["id"])
        assert before["score"] == pytest.approx(19.0)  # верно всё, 19 из 19

        text_answer = next(
            a for a in storage.get_answers(ivanov["id"]) if a["question_id"] == "txt1"
        )
        storage.review_answer(text_answer["id"], accepted=True)
        after = storage.get_attempt(ivanov["id"])
        assert after["score"] == pytest.approx(19.0)
        assert after["percent"] == 100
        assert after["needs_review"] == 0

        # 6б. Сидорова написала «не знаю» — автомат не засчитал. Проверяющий
        #     решает иначе, и балл, процент и оценка меняются вместе.
        sidorova = next(a for a in attempts if a["employee_name"] == "Сидорова А.А.")
        before = storage.get_attempt(sidorova["id"])
        text_answer = next(
            a for a in storage.get_answers(sidorova["id"]) if a["question_id"] == "txt1"
        )
        assert text_answer["score"] == 0.0

        storage.review_answer(text_answer["id"], accepted=True)
        after = storage.get_attempt(sidorova["id"])
        assert after["score"] == pytest.approx(before["score"] + 1.0)
        assert after["percent"] == round(after["score"] / 19 * 100)
        assert after["needs_review"] == 0
        # Оценка пересчитана вместе с процентом, а не осталась прежней.
        assert after["grade_label"] == bundle.test.grading.scale.grade(
            after["percent"]
        ).label

        # 7. Разбор показывает то, что человек видел: порядок и снимки ответов.
        answers_rows = storage.get_answers(ivanov["id"])
        assert [r["shown_index"] for r in answers_rows] == list(range(1, 15))
        assert [r["question_id"] for r in answers_rows] == plans["Иванов И.И."].question_ids
        assert all(r["given_answer"] for r in answers_rows)

        # 8. Отчёт в Excel.
        path = export_attempts(storage, storage.list_attempts(), tmp_path / "отчёт.xlsx")
        book = load_workbook(path)
        assert book.sheetnames == ["Сводная", "По вопросам", "По разделам", "Сложные вопросы"]

        summary = book["Сводная"]
        names = [row[0] for row in summary.iter_rows(min_row=6, values_only=True)]
        assert set(names) == {"Иванов И.И.", "Петров П.П.", "Сидорова А.А."}

        details = list(book["По вопросам"].iter_rows(min_row=2, values_only=True))
        assert len(details) == 42  # 3 попытки × 14 вопросов

        sections = {row[0]: row for row in book["По разделам"].iter_rows(min_row=2, values_only=True)}
        assert set(sections) == {"Охрана труда", "Оборудование", "Вне разделов"}


def _answers_for(bundle: Bundle, plan, name: str) -> dict:
    """Иванов отвечает верно, Петров — наполовину, Сидорова — плохо."""
    answers: dict[str, dict] = {}
    for index, qid in enumerate(plan.question_ids):
        q = bundle.test.question_by_id(qid)
        good = (
            name == "Иванов И.И."
            or (name == "Петров П.П." and index % 2 == 0)
        )
        if q.type is QuestionType.SINGLE:
            option = next(o for o in q.options if o.correct) if good else q.options[1]
            answers[qid] = {"option_id": option.id}
        elif q.type is QuestionType.MULTI:
            ids = [o.id for o in q.options if o.correct] if good else [q.options[2].id]
            answers[qid] = {"option_ids": ids}
        elif q.type is QuestionType.TEXT:
            answers[qid] = {"text": "оформляется наряд-допуск" if good else "не знаю"}
        elif q.type is QuestionType.MATCHING:
            mapping = (
                {p.id: p.id for p in q.pairs}
                if good
                else {q.pairs[0].id: q.pairs[1].id, q.pairs[1].id: q.pairs[0].id}
            )
            answers[qid] = {"pairs": mapping}
        elif q.type is QuestionType.ORDERING:
            order = [i.id for i in q.order]
            answers[qid] = {"order": order if good else list(reversed(order))}
    return answers


def test_by_section_selection_respects_take_count(data_dir, tmp_path):
    """Из каждого раздела берётся ровно столько вопросов, сколько задано."""
    bundle = build_annual_test(tmp_path)
    for seed in range(15):
        plan = build_plan(bundle.test, seed=seed)
        ot = sum(1 for q in plan.question_ids if q.startswith("ot"))
        eq = sum(1 for q in plan.question_ids if q.startswith("eq"))
        assert (ot, eq) == (5, 5)


def test_shuffle_does_not_change_the_score(data_dir, tmp_path):
    """Один и тот же набор ответов даёт один и тот же балл при любом seed."""
    bundle = build_annual_test(tmp_path)
    bundle.test.settings.selection_mode = SelectionMode.ALL

    scores = set()
    for seed in range(10):
        plan = build_plan(bundle.test, seed=seed)
        answers = _answers_for(bundle, plan, "Иванов И.И.")
        scores.add(round(grade_attempt(bundle.test, plan, answers).score, 6))

    assert len(scores) == 1


def test_pinned_option_stays_last_in_real_test(data_dir, tmp_path):
    """«Все перечисленное» не должно всплывать в середине списка."""
    bundle = build_annual_test(tmp_path)
    for seed in range(15):
        plan = build_plan(bundle.test, seed=seed)
        for qid, order in plan.option_order.items():
            if qid.startswith("ot"):
                assert order[-1] == f"{qid}c"
